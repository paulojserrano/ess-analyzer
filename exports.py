"""
exports.py — Excel exports written alongside the HTML report.

Every export reuses the shared analysis helpers so the numbers in Excel are
exactly the numbers in the charts:
  • completions   → analyses._common.triggergo_completions
  • pick events   → analyses.dwell_time.extract_picks  (station-match guard,
                    zero-second picks dropped, ppReady fallback honoured)
  • switch events → analyses.switch_time._release_arrived_gaps, operational
                    swaps only (≤ MAX_OPERATIONAL_SWITCH_S) — longer gaps are
                    starvation and never count as switch time or occupancy.

Public API
----------
save_day_exports(data, cfg, outdir)              per-day workbooks
save_chart_data(registry, outdir)                one sheet of raw rows per chart
save_combined_workbook(completed_days, outdir)   self-auditable all_data.xlsx

Failures are logged (never raised) so one bad export cannot sink a run.
"""
from __future__ import annotations

import logging
import os
import re

import pandas as pd

from analyses._common import triggergo_completions
from analyses.dwell_time import _clipped_occupancy, extract_picks, resolve_pick_start_event
from analyses.switch_time import _release_arrived_gaps
from config import MAX_OPERATIONAL_SWITCH_S, SWITCH_S_FALLBACK, TOTAL_DURATION_COL

_log = logging.getLogger(__name__)

_MAX_CYCLE_S = 7200  # cycle-time outlier bound, matching the charts


def _day_hours(ts: pd.Series) -> pd.DatetimeIndex:
    return pd.date_range(ts.min().normalize(), periods=24, freq="h")


def _operational_switches(lsr: pd.DataFrame, cfg: dict) -> pd.DataFrame:
    gaps = _release_arrived_gaps(lsr, cfg.get("point2ws", {}))
    return gaps[gaps["gap_s"] <= MAX_OPERATIONAL_SWITCH_S].reset_index(drop=True)


def _delivery_legs(tlc: pd.DataFrame) -> tuple[pd.DataFrame, str] | None:
    dest_col = next((c for c in tlc.columns if "目标位置" in str(c)), None)
    leg_col = next((c for c in tlc.columns if "完成耗时" in str(c)), None)
    if not dest_col or not leg_col:
        return None
    df = tlc[tlc[dest_col].astype(str).str.startswith("LABOR")].copy()
    df["_leg"] = pd.to_numeric(df[leg_col], errors="coerce")
    df = df[(df["_leg"] > 0) & (df["_leg"] < 3600)]
    if df.empty:
        return None
    out = pd.DataFrame({
        "Station": df[dest_col].astype(str),
        "Delivery_Leg_s": df["_leg"].astype(float).round(1),
    })
    return out, leg_col


def _cycle_rows(tlc: pd.DataFrame) -> pd.DataFrame | None:
    if TOTAL_DURATION_COL not in tlc.columns:
        return None
    tj = tlc.copy()
    tj[TOTAL_DURATION_COL] = pd.to_numeric(tj[TOTAL_DURATION_COL], errors="coerce")
    tj = tj.dropna(subset=[TOTAL_DURATION_COL])
    return tj[(tj[TOTAL_DURATION_COL] >= 0) & (tj[TOTAL_DURATION_COL] < _MAX_CYCLE_S)]


def _fleet_hourly(lsr: pd.DataFrame) -> pd.DataFrame | None:
    rb_col = next((c for c in lsr.columns if "机器人编号" in str(c)), None)
    rbt_col = next((c for c in lsr.columns if "机器人类型" in str(c)), None)
    if not rb_col or not rbt_col:
        return None
    d = lsr.copy()
    d["ts"] = pd.to_datetime(d["时间戳"], errors="coerce")
    d = d.dropna(subset=["ts", rb_col, rbt_col])
    rows: list[dict] = []
    for rtype, fleet in d.groupby(rbt_col):
        fleet_size = fleet[rb_col].nunique()
        for hr, cnt in fleet.groupby(fleet["ts"].dt.floor("h"))[rb_col].nunique().items():
            rows.append({
                "Hour": hr.strftime("%H:00"),
                "Robot_Type": rtype,
                "Robots_Seen": int(cnt),
                "Fleet_Size": fleet_size,
                "Seen_pct": round(cnt / fleet_size * 100, 1) if fleet_size else None,
            })
    return pd.DataFrame(rows) if rows else None


def _fmt_hours(p: pd.DataFrame) -> pd.DataFrame:
    out = p.copy()
    out.columns = [c.strftime("%H:00") for c in out.columns]
    out.index.name = "Station"
    return out.round(1)


# ── per-day workbooks ─────────────────────────────────────────────────────────

def save_day_exports(data: dict, cfg: dict, outdir: str) -> list[str]:
    """Write the per-day Excel workbooks.  Returns the paths written."""
    written: list[str] = []
    lsr = data.get("station")
    tlc = data.get("lifecycle")
    ws_order = cfg.get("ws_order", [])

    def _attempt(name: str, fn) -> None:
        try:
            path = fn(os.path.join(outdir, name))
            if path:
                written.append(path)
        except Exception as exc:  # one failing export must not stop the rest
            _log.warning("Excel export %s skipped: %s: %s", name, type(exc).__name__, exc)

    if lsr is not None and ws_order:
        def _throughput(path: str) -> str | None:
            tgo = triggergo_completions(lsr, cfg)
            if tgo.empty:
                return None
            pivot = (
                tgo.groupby([tgo["ts"].dt.floor("h"), "station"]).size()
                .unstack(fill_value=0)
            )
            pivot = pivot[[w for w in ws_order if w in pivot.columns]]
            out = pivot.copy()
            out.index = [h.strftime("%H:00") for h in out.index]
            out.index.name = "Hour"
            out["TOTAL"] = out.sum(axis=1)
            out.loc["TOTAL"] = out.sum()
            with pd.ExcelWriter(path) as w:
                out.to_excel(w, sheet_name="throughput_by_hour")
                tgo[["时间戳", "station", "事件类型"]].rename(columns={
                    "时间戳": "Timestamp", "station": "Station", "事件类型": "Event",
                }).to_excel(w, sheet_name="raw_events", index=False)
            return path
        _attempt("throughput_by_workstation_hour.xlsx", _throughput)

        def _dwell(path: str) -> str | None:
            picks = extract_picks(lsr, cfg)
            picks = picks[picks["station"].isin(ws_order)]
            if picks.empty:
                return None
            hrs = _day_hours(picks["hour_dt"])
            avg_pick = (
                picks.groupby(["station", "hour_dt"])["pick_s"].mean()
                .unstack().reindex(index=ws_order, columns=hrs)
            )
            pick_occ = _clipped_occupancy(picks, "arr_ts", "tg_ts", ws_order, hrs)
            tgo = triggergo_completions(lsr, cfg)
            actual = (
                tgo.groupby(["station", tgo["ts"].dt.floor("h")]).size()
                .unstack(fill_value=0)
                .reindex(index=ws_order, columns=hrs, fill_value=0)
            )
            sw = _operational_switches(lsr, cfg)
            with pd.ExcelWriter(path) as w:
                _fmt_hours(avg_pick).to_excel(w, sheet_name="avg_pick_time_s")
                _fmt_hours(actual).to_excel(w, sheet_name="actual_completions")
                _fmt_hours(pick_occ).to_excel(w, sheet_name="pick_seconds_per_hour")
                _fmt_hours(pick_occ / 36.0).to_excel(w, sheet_name="pick_occupancy_pct")
                if not sw.empty:
                    sw_occ = _clipped_occupancy(sw, "rel_ts", "arr_ts", ws_order, hrs)
                    full = (pick_occ + sw_occ).clip(upper=3600.0)
                    _fmt_hours(sw_occ).to_excel(w, sheet_name="switch_seconds_per_hour")
                    _fmt_hours(full / 36.0).to_excel(w, sheet_name="station_occupancy_pct")
                    sw.assign(Hour=sw["rel_ts"].dt.strftime("%H:00")).rename(columns={
                        "station": "Station", "gap_s": "Switch Time (s)",
                        "rel_ts": "Release", "arr_ts": "Next Arrived",
                    }).to_excel(w, sheet_name="raw_switch_events", index=False)
                picks.assign(Hour=picks["hour_dt"].dt.strftime("%H:00"))[
                    ["robot", "station", "Hour", "arr_ts", "tg_ts", "pick_s"]
                ].rename(columns={
                    "robot": "Robot", "station": "Station", "arr_ts": "Start",
                    "tg_ts": "TriggerGo", "pick_s": "Pick Time (s)",
                }).to_excel(w, sheet_name="raw_pick_events", index=False)
            return path
        _attempt("dwell_capacity_utilisation.xlsx", _dwell)

    if lsr is not None:
        def _fleet(path: str) -> str | None:
            df = _fleet_hourly(lsr)
            if df is None:
                return None
            with pd.ExcelWriter(path) as w:
                df.to_excel(w, sheet_name="robots_seen_by_hour", index=False)
            return path
        _attempt("fleet_utilization_profile.xlsx", _fleet)

    if tlc is not None:
        def _cycle(path: str) -> str | None:
            tj = _cycle_rows(tlc)
            if tj is None or tj.empty:
                return None
            cyc = (tj[TOTAL_DURATION_COL] / 60).rename("cycle_min")
            with pd.ExcelWriter(path) as w:
                pd.concat([tj[TOTAL_DURATION_COL], cyc], axis=1).to_excel(
                    w, sheet_name="cycle_time_intervals", index=False)
                cyc.describe(percentiles=[0.5, 0.75, 0.9, 0.95, 0.99]).round(2) \
                    .reset_index().to_excel(w, sheet_name="summary_stats", index=False)
                tj.to_excel(w, sheet_name="raw_lifecycle", index=False)
            return path
        _attempt("cycle_time_distribution.xlsx", _cycle)

        def _legs(path: str) -> str | None:
            res = _delivery_legs(tlc)
            if res is None:
                return None
            legs, _ = res
            stats = (
                legs.groupby("Station")["Delivery_Leg_s"]
                .describe(percentiles=[0.25, 0.5, 0.75, 0.9]).round(1)
            )
            order = [w for w in ws_order if w in stats.index] + \
                    [s for s in stats.index if s not in ws_order]
            with pd.ExcelWriter(path) as w:
                stats.reindex(order).to_excel(w, sheet_name="delivery_leg_stats_by_station")
                legs.to_excel(w, sheet_name="raw_deliveries", index=False)
            return path
        _attempt("fleet_delivery_leg.xlsx", _legs)

    return written


# ── per-chart raw data ────────────────────────────────────────────────────────

def _sheet_name(chart_id: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "_", chart_id)[:31] or "chart"
    name, n = base, 2
    while name in used:
        suffix = f"_{n}"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def save_chart_data(registry: list[dict], outdir: str,
                    filename: str = "chart_data.xlsx") -> str | None:
    """One sheet per chart with the tabular rows behind it (ChartResult
    'raw_data' → 'rows'), plus an index sheet with title/source/method."""
    sheets: list[tuple[dict, pd.DataFrame]] = []
    for entry in registry:
        rows = (entry.get("raw_data") or {}).get("rows")
        if rows:
            try:
                sheets.append((entry, pd.DataFrame(rows)))
            except Exception as exc:
                _log.warning("Chart data for %s skipped: %s", entry.get("id"), exc)
    if not sheets:
        return None
    path = os.path.join(outdir, filename)
    try:
        used: set[str] = {"_Index"}
        index_rows = []
        with pd.ExcelWriter(path) as w:
            named = [(_sheet_name(e["id"], used), e, df) for e, df in sheets]
            for name, e, _ in named:
                index_rows.append({"Sheet": name, "Chart": e.get("title", ""),
                                   "Source": e.get("source", ""), "Method": e.get("method", "")})
            pd.DataFrame(index_rows).to_excel(w, sheet_name="_Index", index=False)
            for name, _, df in named:
                df.to_excel(w, sheet_name=name, index=False)
        return path
    except Exception as exc:
        _log.warning("chart_data.xlsx skipped: %s: %s", type(exc).__name__, exc)
        return None


# ── combined, self-auditable workbook ─────────────────────────────────────────

def save_combined_workbook(completed_days: list[dict], base_outdir: str) -> str | None:
    """Write all_data.xlsx: raw data + formula-derived + computed sheets.

    Sheet taxonomy (per day, prefixed D{n}_ when multi-day):
      *_raw           — verbatim source data, no filtering
      tp_events       — completion events (AMR triggerGo at a station) + Hour
      tp_pivot        — throughput matrix using live COUNTIFS (formula)
      pick_events     — pick pairs (shared pairing) with Pick_s formula column
      switch_events   — operational release→arrived swaps with Switch_s formula
      dwell_derived   — utilisation formula chain, one row per Station×Hour
      cycle_events    — lifecycle rows within bounds + Cycle_min formula column
      cycle_stats     — descriptive stats (computed)
      delivery_leg    — delivery-leg durations + Station column
      leg_stats       — per-station percentiles (computed)
      fleet_util      — robots seen per hour (computed)
    """
    try:
        return _write_combined(completed_days, base_outdir)
    except Exception as exc:
        _log.warning("all_data.xlsx skipped: %s: %s", type(exc).__name__, exc)
        return None


def _write_combined(completed_days: list[dict], base_outdir: str) -> str:
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    multi = len(completed_days) > 1
    out_path = os.path.join(base_outdir, "all_data.xlsx")

    dark, blue, lblue, amber = "0F172A", "1D4ED8", "EFF6FF", "FEF3C7"

    def style_header(ws) -> None:
        for col in range(1, (ws.max_column or 1) + 1):
            c = ws.cell(1, col)
            if c.value is not None:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor=dark)

    def note(ws, row, col, text) -> None:
        c = ws.cell(row, col, text)
        c.font = Font(italic=True, color=blue)
        c.fill = PatternFill("solid", fgColor=lblue)

    def method(ws, row, col, text) -> None:
        c = ws.cell(row, col, text)
        c.font = Font(italic=True, color="92400E")
        c.fill = PatternFill("solid", fgColor=amber)

    def header_cell(ws, col, text) -> None:
        c = ws.cell(1, col, text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=dark)

    def sn(base: str, prefix: str) -> str:
        return (f"{prefix}{base}" if prefix else base)[:31]

    def section(name: str, fn) -> None:
        try:
            fn()
        except Exception as exc:
            _log.warning("all_data.xlsx: sheet %s skipped: %s: %s",
                         name, type(exc).__name__, exc)

    with pd.ExcelWriter(out_path, engine="openpyxl") as w:
        wb = w.book
        legend = pd.DataFrame([
            ["cb_raw", "Raw", "Complete callback sheet — no filtering"],
            ["st_raw", "Raw", "Complete station sheet — all labour-station robot events"],
            ["lc_raw", "Raw", "Complete lifecycle sheet — full container journey records"],
            ["tp_events", "Data", "Completions: delivery-AMR triggerGo events at mapped LABOR stations, with Hour bucket"],
            ["tp_pivot", "Formula", "Throughput matrix — cells are live COUNTIFS(tp_events)"],
            ["pick_events", "Data", "Pick pairs (start event → triggerGo, same station); Pick_s is an Excel formula"],
            ["switch_events", "Data", f"Operational swaps only: release → next arrived ≤ {MAX_OPERATIONAL_SWITCH_S:.0f} s; Switch_s is an Excel formula"],
            ["dwell_derived", "Formula", "Utilisation chain: avg pick, avg switch, implied TPH, util% — live formulas"],
            ["cycle_events", "Data", "Lifecycle rows within 0–2 h; Cycle_min is an Excel formula"],
            ["cycle_stats", "Computed", "Cycle-time descriptive statistics"],
            ["delivery_leg", "Data", "Delivery legs to LABOR stations from the lifecycle sheet"],
            ["leg_stats", "Computed", "Per-station delivery-leg percentiles"],
            ["fleet_util", "Computed", "Robots seen per hour ÷ fleet size (activity, not concurrency)"],
        ], columns=["Sheet suffix", "Type", "Description"])
        legend.to_excel(w, sheet_name="_Index", index=False)
        ix = wb["_Index"]
        style_header(ix)
        ix.column_dimensions["A"].width = 18
        ix.column_dimensions["B"].width = 12
        ix.column_dimensions["C"].width = 90
        if multi:
            c = ix.cell(len(legend) + 3, 1,
                        "Multi-day run: each sheet is prefixed D1_, D2_, … — " +
                        ", ".join(f"D{i + 1} = {d['label']}" for i, d in enumerate(completed_days)))
            c.font = Font(italic=True, color=blue)

        for day_idx, day in enumerate(completed_days):
            data, cfg = day["data"], day["cfg"]
            p = f"D{day_idx + 1}_" if multi else ""
            cb, lsr, tlc = data.get("callback"), data.get("station"), data.get("lifecycle")
            ws_order = cfg.get("ws_order", [])
            names: dict[str, str | None] = {"tp": None, "pick": None, "switch": None}

            def raw_sheets() -> None:
                for base, df in (("cb_raw", cb), ("st_raw", lsr), ("lc_raw", tlc)):
                    if df is not None:
                        df.to_excel(w, sheet_name=sn(base, p), index=False)
                        style_header(wb[sn(base, p)])
            section("raw", raw_sheets)

            def tp_events() -> None:
                if lsr is None or not ws_order:
                    return
                tgo = triggergo_completions(lsr, cfg).copy()
                tgo["Hour"] = tgo["ts"].dt.strftime("%H:00")
                name = sn("tp_events", p)
                tgo[["时间戳", "Hour", "station"]].rename(
                    columns={"时间戳": "Timestamp", "station": "Station"}
                ).to_excel(w, sheet_name=name, index=False)
                style_header(wb[name])
                note(wb[name], 1, 4,
                     "Hour = Timestamp floored to HH:00 | Station = LABOR station mapped from 位置编号 | "
                     "Filter: 事件类型 == 'triggerGo', delivery AMR only, at mapped stations")
                names["tp"] = name
            section("tp_events", tp_events)

            def tp_pivot() -> None:
                tp = names["tp"]
                if not tp or not ws_order:
                    return
                pws = wb.create_sheet(sn("tp_pivot", p))
                n_st = len(ws_order)
                hours = [f"{h:02d}:00" for h in range(24)]
                total_col = n_st + 2
                pws.cell(1, 1, "Hour")
                for ci, st in enumerate(ws_order, 2):
                    pws.cell(1, ci, st)
                pws.cell(1, total_col, "TOTAL")
                style_header(pws)
                note(pws, 2, 1, "FORMULA →")
                note(pws, 2, 2, f"=COUNTIFS('{tp}'!$C:$C,[Station],'{tp}'!$B:$B,[Hour])")
                for ri, hr in enumerate(hours, 3):
                    pws.cell(ri, 1, hr)
                    for ci in range(2, n_st + 2):
                        col = get_column_letter(ci)
                        pws.cell(ri, ci, f"=COUNTIFS('{tp}'!$C:$C,{col}$1,'{tp}'!$B:$B,$A{ri})")
                    pws.cell(ri, total_col, f"=SUM(B{ri}:{get_column_letter(n_st + 1)}{ri})")
                tr = 3 + len(hours)
                pws.cell(tr, 1, "TOTAL")
                for ci in range(2, total_col + 1):
                    col = get_column_letter(ci)
                    pws.cell(tr, ci, f"=SUM({col}3:{col}{tr - 1})").font = Font(bold=True)
            section("tp_pivot", tp_pivot)

            def pick_events() -> None:
                if lsr is None:
                    return
                picks = extract_picks(lsr, cfg)
                if picks.empty:
                    return
                start_lbl = "Start_" + resolve_pick_start_event(cfg)
                df = pd.DataFrame({
                    "Robot_ID": picks["robot"].values,
                    "Station": picks["station"].astype(str).values,
                    "Hour": picks["hour_dt"].dt.strftime("%H:00").values,
                    start_lbl: picks["arr_ts"].values,
                    "TriggerGo": picks["tg_ts"].values,
                })
                name = sn("pick_events", p)
                df.to_excel(w, sheet_name=name, index=False)
                ews = wb[name]
                style_header(ews)
                header_cell(ews, 6, "Pick_s")
                note(ews, 1, 7,
                     f"Pick_s = (TriggerGo − {start_lbl}) × 86400 | pairing shared with the "
                     "charts: same station, 0 < Pick_s < 3600")
                for r in range(2, len(df) + 2):
                    ews.cell(r, 6, f"=(E{r}-D{r})*86400")
                names["pick"] = name
            section("pick_events", pick_events)

            def switch_events() -> None:
                if lsr is None:
                    return
                sw = _operational_switches(lsr, cfg)
                if sw.empty:
                    return
                df = pd.DataFrame({
                    "Station": sw["station"].values,
                    "Hour": sw["rel_ts"].dt.strftime("%H:00").values,
                    "Release_ts": sw["rel_ts"].values,
                    "Arrived_ts": sw["arr_ts"].values,
                })
                name = sn("switch_events", p)
                df.to_excel(w, sheet_name=name, index=False)
                sws = wb[name]
                style_header(sws)
                header_cell(sws, 5, "Switch_s")
                note(sws, 1, 6,
                     "Switch_s = (Arrived_ts − Release_ts) × 86400 | release → next arrived at the "
                     f"same station | operational swaps only (≤ {MAX_OPERATIONAL_SWITCH_S:.0f} s); "
                     "longer gaps are starvation, see the Station Starvation chart")
                for r in range(2, len(df) + 2):
                    sws.cell(r, 5, f"=(D{r}-C{r})*86400")
                names["switch"] = name
            section("switch_events", switch_events)

            def dwell_derived() -> None:
                tp, pk, swn = names["tp"], names["pick"], names["switch"]
                if not (tp and pk and ws_order):
                    return
                sw_fixed = float(cfg.get("switch_s_fixed", SWITCH_S_FALLBACK))
                dws = wb.create_sheet(sn("dwell_derived", p))
                headers = ["Station", "Hour", "Avg_Pick_s", "Avg_Switch_s",
                           "Implied_TPH_fixed", "Implied_TPH_actual", "Actual_Completions",
                           "Util_fixed_%", "Util_actual_%"]
                for ci, h in enumerate(headers, 1):
                    dws.cell(1, ci, h)
                style_header(dws)
                legend_row = [
                    "(row value)", "(row value)",
                    "AVERAGEIFS(pick_events[Pick_s], Station, Hour)",
                    "AVERAGEIFS(switch_events[Switch_s], Station, Hour)",
                    f"3600 / (Avg_Pick_s + {sw_fixed:g})  ← user-set fixed switch",
                    "3600 / (Avg_Pick_s + Avg_Switch_s)",
                    "COUNTIFS(tp_events[Station], Station, tp_events[Hour], Hour)",
                    "Actual_Completions / Implied_TPH_fixed × 100",
                    "Actual_Completions / Implied_TPH_actual × 100",
                ]
                for ci, txt in enumerate(legend_row, 1):
                    note(dws, 2, ci, txt)
                row = 3
                for st in ws_order:
                    for h in range(24):
                        hr = f"{h:02d}:00"
                        dws.cell(row, 1, st)
                        dws.cell(row, 2, hr)
                        dws.cell(row, 3, f"=IFERROR(AVERAGEIFS('{pk}'!$F:$F,'{pk}'!$B:$B,$A{row},"
                                         f"'{pk}'!$C:$C,$B{row}),\"\")")
                        if swn:
                            dws.cell(row, 4, f"=IFERROR(AVERAGEIFS('{swn}'!$E:$E,'{swn}'!$A:$A,$A{row},"
                                             f"'{swn}'!$B:$B,$B{row}),\"\")")
                            dws.cell(row, 6, f"=IFERROR(3600/(C{row}+D{row}),\"\")")
                            dws.cell(row, 9, f"=IFERROR(G{row}/F{row}*100,\"\")")
                        dws.cell(row, 5, f"=IFERROR(3600/(C{row}+{sw_fixed:g}),\"\")")
                        dws.cell(row, 7, f"=COUNTIFS('{tp}'!$C:$C,$A{row},'{tp}'!$B:$B,$B{row})")
                        dws.cell(row, 8, f"=IFERROR(G{row}/E{row}*100,\"\")")
                        row += 1
            section("dwell_derived", dwell_derived)

            def cycle() -> None:
                if tlc is None:
                    return
                tj = _cycle_rows(tlc)
                if tj is None or tj.empty:
                    return
                name = sn("cycle_events", p)
                tj.to_excel(w, sheet_name=name, index=False)
                cws = wb[name]
                style_header(cws)
                dur_col = get_column_letter(list(tj.columns).index(TOTAL_DURATION_COL) + 1)
                nxt = len(tj.columns) + 1
                header_cell(cws, nxt, "Cycle_min")
                note(cws, 1, nxt + 1, f"Cycle_min = {TOTAL_DURATION_COL} ÷ 60 | "
                                      "outliers filtered: 0 ≤ duration < 7200 s")
                for r in range(2, len(tj) + 2):
                    cws.cell(r, nxt, f"={dur_col}{r}/60")
                stats = (tj[TOTAL_DURATION_COL].div(60)
                         .describe(percentiles=[0.5, 0.75, 0.9, 0.95, 0.99])
                         .round(2).reset_index())
                stats.columns = ["Metric", "Cycle_min"]
                sname = sn("cycle_stats", p)
                stats.to_excel(w, sheet_name=sname, index=False)
                style_header(wb[sname])
                method(wb[sname], 1, 3, "Computed: percentiles of (total duration ÷ 60); "
                                         "rows ≥ 7200 s or blank are excluded.")
            section("cycle", cycle)

            def legs() -> None:
                if tlc is None:
                    return
                res = _delivery_legs(tlc)
                if res is None:
                    return
                df, leg_col = res
                name = sn("delivery_leg", p)
                df.to_excel(w, sheet_name=name, index=False)
                style_header(wb[name])
                note(wb[name], 1, 3, f"Delivery_Leg_s = '{leg_col}' | destination starts with "
                                     "'LABOR', 0 < value < 3600")
                stats = (df.groupby("Station")["Delivery_Leg_s"]
                         .describe(percentiles=[0.25, 0.5, 0.75, 0.9]).round(1))
                sname = sn("leg_stats", p)
                stats.to_excel(w, sheet_name=sname)
                style_header(wb[sname])
                method(wb[sname], 1, len(stats.columns) + 2,
                       "Computed: describe() of Delivery_Leg_s per Station.")
            section("legs", legs)

            def fleet() -> None:
                if lsr is None:
                    return
                df = _fleet_hourly(lsr)
                if df is None:
                    return
                name = sn("fleet_util", p)
                df.to_excel(w, sheet_name=name, index=False)
                style_header(wb[name])
                method(wb[name], 1, len(df.columns) + 2,
                       "Computed: Robots_Seen = unique robot IDs with any station event in the "
                       "hour; Fleet_Size = unique robots of that type all day. This is an "
                       "activity measure — the concurrent on-task profile is in the HTML report.")
            section("fleet", fleet)

        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
    return out_path


__all__ = ["save_day_exports", "save_chart_data", "save_combined_workbook"]
